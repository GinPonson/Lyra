"""Schwab XLSX 解析测试。

**三类结构被显式拒绝**，各自防一种静默错误：

| 结构 | 不拒绝会怎样 |
| --- | --- |
| 公式 | 导入的是公式而非其值；Excel 重算后金额可能与账本不一致 |
| 合并单元格 | 值只在左上角，其余格为空——行会缺列 |
| 隐藏工作表 | 用户看不到被导入的数据 |

**多个候选工作表时报歧义而不是挑第一个**——挑错了会让用户在一个他以为正确的文件上
得到错误的账本数据。
"""

from __future__ import annotations

import io
from datetime import datetime

import openpyxl
import pytest

from lyra.domain.ledger.errors import IngestError
from lyra.infrastructure.imports.schwab_table import SCHWAB_TRANSACTION_HEADERS
from lyra.infrastructure.imports.schwab_xlsx import parse_schwab_xlsx

FILE_NAME = "Individual_XXX276_Transactions.xlsx"


def _workbook(*, title: str = "Transactions") -> openpyxl.Workbook:
    """建立工作簿。"""
    workbook = openpyxl.Workbook()
    workbook.active.title = title
    return workbook


def _transaction_sheet(workbook, *, rows=None, title: str = "Transactions"):
    """写入表头与数据行。"""
    sheet = workbook.create_sheet(title) if title not in workbook.sheetnames else workbook[title]
    sheet.append(list(SCHWAB_TRANSACTION_HEADERS))
    for row in rows or (["01/05/2026", "Buy", "INTC", "INTEL CORP", "10", "$40.5", "", "-$405"],):
        sheet.append(row)
    return sheet


def _content(workbook) -> bytes:
    """序列化工作簿。"""
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _parse(content: bytes, **overrides):
    """跑一次解析。"""
    return parse_schwab_xlsx(
        content, file_name=overrides.pop("file_name", FILE_NAME),
        source_hash="h", source_reference="r", **overrides,
    )


# --------------------------------------------------------------------------
# 正常
# --------------------------------------------------------------------------


def test_parses_a_valid_workbook() -> None:
    """合法工作簿解析出一行交易。"""
    workbook = _workbook()
    _transaction_sheet(workbook)

    batch = _parse(_content(workbook))

    assert batch.source_type == "excel"
    assert batch.account == "XXX276"
    assert len(batch.rows) == 1
    assert batch.rows[0].symbol == "INTC"


def test_source_type_is_excel_not_xlsx() -> None:
    """**来源类型是 ``excel``**（schema 的 CHECK 只允许这个值）。"""
    workbook = _workbook()
    _transaction_sheet(workbook)

    assert _parse(_content(workbook)).source_type == "excel"


def test_explicit_source_type_can_be_overridden() -> None:
    """调用方可显式指定来源类型。"""
    workbook = _workbook()
    _transaction_sheet(workbook)

    assert _parse(_content(workbook), source_type="csv").source_type == "csv"


def test_header_need_not_be_the_first_row_of_the_sheet() -> None:
    """表头前可以有完全空行。"""
    workbook = _workbook()
    sheet = workbook.active
    sheet.title = "Transactions"
    sheet.append([])
    sheet.append([])
    sheet.append(list(SCHWAB_TRANSACTION_HEADERS))
    sheet.append(["01/05/2026", "Buy", "INTC", "INTEL CORP", "10", "$40.5", "", "-$405"])

    batch = _parse(_content(workbook))

    assert len(batch.rows) == 1


def test_empty_rows_between_data_are_skipped() -> None:
    """数据行之间的空行被跳过。"""
    workbook = _workbook()
    sheet = _transaction_sheet(workbook)
    sheet.append([])
    sheet.append(["01/06/2026", "Buy", "AMD", "ADVANCED MICRO", "5", "$100", "", "-$500"])

    batch = _parse(_content(workbook))

    assert len(batch.rows) == 2


def test_extra_trailing_empty_columns_are_trimmed() -> None:
    """尾部多余空列被裁掉，仍然解析成功。"""
    workbook = _workbook()
    sheet = _transaction_sheet(workbook)
    # 在第 10 列写一个空值，制造"多出的列"。
    sheet.cell(row=2, column=10).value = None

    assert len(_parse(_content(workbook)).rows) == 1


# --------------------------------------------------------------------------
# 工作表选择
# --------------------------------------------------------------------------


def test_sheet_whose_first_non_empty_row_is_not_the_header_is_skipped() -> None:
    """**表头不在首个非空行就判定该表不是交易表**，不继续往下找。

    继续找会在一份"前面有标题行、后面才是数据"的表里选中它，而那些标题行与交易无关。
    """
    workbook = _workbook()
    sheet = workbook.active
    sheet.title = "Transactions"
    sheet.append(["My Portfolio Summary"])
    sheet.append(list(SCHWAB_TRANSACTION_HEADERS))
    sheet.append(["01/05/2026", "Buy", "INTC", "INTEL CORP", "10", "$40.5", "", "-$405"])

    with pytest.raises(IngestError, match="没有 Schwab Transactions 表"):
        _parse(_content(workbook))


def test_hidden_transaction_sheet_is_rejected() -> None:
    """**交易工作表不可隐藏**——用户看不到被导入的数据。"""
    workbook = _workbook()
    sheet = _transaction_sheet(workbook)
    sheet.sheet_state = "hidden"
    # 加一个可见的空表，让"没有候选"不是失败原因。
    workbook.create_sheet("Notes")

    with pytest.raises(IngestError, match="不可隐藏"):
        _parse(_content(workbook))


def test_multiple_candidates_are_ambiguous() -> None:
    """**多个候选取歧义**，并列出全部候选工作表名。"""
    workbook = _workbook()
    _transaction_sheet(workbook, title="Transactions")
    _transaction_sheet(workbook, title="Transactions2")

    with pytest.raises(IngestError) as raised:
        _parse(_content(workbook))

    message = str(raised.value)
    assert "多个" in message
    assert "Transactions" in message and "Transactions2" in message


def test_workbook_without_a_transaction_sheet_is_rejected() -> None:
    """没有候选表时报错。"""
    workbook = _workbook()
    workbook.active.append(["unrelated"])

    with pytest.raises(IngestError, match="没有 Schwab Transactions 表"):
        _parse(_content(workbook))


# --------------------------------------------------------------------------
# 单元格结构
# --------------------------------------------------------------------------


def test_formula_cells_are_rejected() -> None:
    """**公式单元格拒绝**：导入公式本身没有意义，而它的值会随 Excel 重算而变。"""
    workbook = _workbook()
    sheet = _transaction_sheet(workbook)
    sheet["E2"] = "=10*1"

    with pytest.raises(IngestError, match="不允许公式或合并单元格"):
        _parse(_content(workbook))


def test_unreadable_content_is_rejected() -> None:
    """**不是工作簿的字节报错**（加密、损坏或格式不符）。"""
    with pytest.raises(IngestError, match="无法读取或已加密"):
        _parse(b"this is not a workbook")


def test_date_cells_are_rendered_month_day_year() -> None:
    """**日期单元格渲染为 ``MM/DD/YYYY``。**

    Excel 单元格里的日期是**朴素时间**（Excel 本身不存时区），因此这里写入朴素值。
    实现按 UTC 解释它——一份 Schwab 导出里的日期就是那一天的日期，不做时区换算。
    """
    workbook = _workbook()
    sheet = _transaction_sheet(workbook)
    sheet["A2"] = datetime(2026, 1, 5)

    batch = _parse(_content(workbook))

    assert batch.rows[0].date == "01/05/2026"


def test_boolean_cells_are_rejected() -> None:
    """布尔既不是文本也不是金额；当成 "True" 会污染字段。"""
    workbook = _workbook()
    sheet = _transaction_sheet(workbook)
    sheet["E2"] = True

    with pytest.raises(IngestError, match="单元格类型不支持"):
        _parse(_content(workbook))


def test_blank_trailing_cells_are_read_as_empty_strings() -> None:
    """**尾部单元格为空时按空串处理，不报列数错误。**

    这是与参考实现一致的行为：读取时按 `max(8, 实际列数)` 取格，因此即使最后两格为空，
    行仍有八格（值为空串），映射随之成功——费用与净额为空的记录本身是合法的。

    我在写这条测试时预期它会报"列数不足"，实际不会。**实现是对的。**
    """
    workbook = _workbook()
    sheet = _transaction_sheet(workbook)
    sheet.cell(row=2, column=7).value = None
    sheet.cell(row=2, column=8).value = None

    batch = _parse(_content(workbook))

    assert batch.rows[0].fees == ""
    assert batch.rows[0].amount == ""


def test_row_with_too_few_columns_is_rejected() -> None:
    """**真正列数不足时拒绝**：让表头只有四列，数据行也随之只有四列。

    此时表头本身就不是八列，因此该表不会被识别为交易表——错误来自"没有交易表"，
    而不是"列数不足"。这条测试记录的是**两条判定路径的先后顺序**。
    """
    workbook = _workbook()
    sheet = workbook.active
    sheet.title = "Transactions"
    sheet.append(list(SCHWAB_TRANSACTION_HEADERS[:4]))
    sheet.append(["01/05/2026", "Buy", "INTC", "INTEL CORP"])

    with pytest.raises(IngestError, match="没有 Schwab Transactions 表"):
        _parse(_content(workbook))


# --------------------------------------------------------------------------
# 与 CSV 的一致性
# --------------------------------------------------------------------------


def test_xlsx_rows_equal_csv_rows_for_the_same_data() -> None:
    """**同一份数据经 xlsx 与经 csv 产出相同的领域字段。**

    这是跨容器一致性的真正依据：两者都走 `schwab_table` 的规则，因此字段不可能不同。
    """
    from lyra.infrastructure.imports.schwab_csv import parse_schwab_csv

    workbook = _workbook()
    _transaction_sheet(workbook)
    from_xlsx = _parse(_content(workbook))

    csv_text = (
        b"Date,Action,Symbol,Description,Quantity,Price,Fees & Comm,Amount\n"
        b"01/05/2026,Buy,INTC,INTEL CORP,10,$40.5,,-$405\n"
    )
    from_csv = parse_schwab_csv(
        csv_text, file_name=FILE_NAME, source_hash="h", source_reference="r"
    )

    assert from_xlsx.rows == from_csv.rows
    assert from_xlsx.account == from_csv.account
    # 来源类型不同是预期的——那正是用来区分数据来自哪种容器的字段。
    assert from_xlsx.source_type == "excel"
    assert from_csv.source_type == "csv"
